from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parents[1]
source = root / 'docs/洞见_AI游戏电竞赛道_游戏及AI应用说明.docx'
output = source.with_name(source.stem + '_仿宋版.docx')
doc = Document(source)
before = [p.text for p in doc.paragraphs], [[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]

def uniform(rpr):
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for key in list(fonts.attrib):
        del fonts.attrib[key]
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + key), '仿宋')
    for key in ('b', 'bCs', 'i', 'iCs'):
        nodes = rpr.findall(qn('w:' + key))
        for node in nodes:
            rpr.remove(node)
        node = OxmlElement('w:' + key)
        node.set(qn('w:val'), '0')
        rpr.append(node)

for style in doc.styles:
    uniform(style.element.get_or_add_rPr())
for part in doc.part.package.parts:
    element = getattr(part, 'element', None)
    if element is None:
        continue
    for rpr in list(element.iter(qn('w:rPr'))):
        uniform(rpr)
    for run in list(element.iter(qn('w:r'))):
        rpr = run.find(qn('w:rPr'))
        if rpr is None:
            rpr = OxmlElement('w:rPr')
            run.insert(0, rpr)
        uniform(rpr)
doc.save(output)
check = Document(output)
after = [p.text for p in check.paragraphs], [[[c.text for c in r.cells] for r in t.rows] for t in check.tables]
assert before == after, 'Content changed unexpectedly'
with ZipFile(output) as z:
    assert z.testzip() is None
for element in (check.element, check.styles.element):
    for fonts in element.iter(qn('w:rFonts')):
        assert all(fonts.get(qn('w:'+k)) == '仿宋' for k in ('ascii','hAnsi','eastAsia','cs'))
    for key in ('b','bCs'):
        assert all(n.get(qn('w:val')) == '0' for n in element.iter(qn('w:'+key)))
print('Verified: text and tables preserved; all fonts FangSong; bold disabled.')
print(str(output))
