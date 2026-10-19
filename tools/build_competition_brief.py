from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.72)
s.left_margin = s.right_margin = Inches(.8)
for name, size in [('Normal',11),('Title',23),('Subtitle',12),('Heading 1',15),('Heading 2',12)]:
    st=doc.styles[name]
    st.font.name='Microsoft YaHei'
    st.font.size=Pt(size)
    st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.line_spacing=1.25
doc.styles['Normal'].paragraph_format.widow_control=True

def p(text): doc.add_paragraph(text)
def h(text): doc.add_heading(text,1)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for i,values in enumerate([headers]+rows):
        row=t.rows[0] if i==0 else t.add_row()
        pr=row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
        if i==0: pr.append(OxmlElement('w:tblHeader'))
        for cell,text,w in zip(row.cells,values,widths):
            cell.width=Inches(w); cell.text=text
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E7EDF1' if i==0 else ('F6F7F8' if i%2==0 else 'FFFFFF')); cp.append(sh)
            edges=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); edges.append(e)
            cp.append(edges)
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'100'); e.set(qn('w:type'),'dxa'); mar.append(e)
            cp.append(mar)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(2); par.paragraph_format.line_spacing=1.15
                for r in par.runs: r.font.size=Pt(10); r.bold=i==0
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

doc.add_paragraph('洞见游戏说明及 AI 应用说明','Title')
doc.add_paragraph('AI 游戏电竞赛道参赛作品说明','Subtitle')
p('作品类型：2D 横版剧情冒险 Demo\n运行平台：Windows\n开发引擎：Godot 4.7　主要语言：GDScript')
p('《洞见》以重庆城市历史与人文记忆为主题，通过五个玩法不同的章节，呈现现代青年从匆忙打卡到理解城市与他人的转变。AI 深度参与了项目框架、美术内容、玩法脚本、剧情迭代及测试验证。本说明介绍作品体验、AI 工具分工、技术实现与可核验的创作过程。')
h('一 作品故事与创作主题')
p('玩家扮演急着前往洪崖洞打卡的现代青年陈默。他为了赶路选择抄近道，却仍遇上传单阿姨，在烦躁与追逐中触碰石扣，意外进入重庆的历史残影。玩家将走进磁器口窑场、中山古镇与战时防空洞，在劳动、承诺和救助中逐渐理解这座城市，最后回到当代洪崖洞。')
p('作品以“从匆忙打卡到真正看见”为叙事主线，将吊脚楼、石阶、窑火、古镇商贸与防空洞融入可操作的关卡。玩家通过行动体会故事，而人物成长则体现在面对相同地点时作出的不同选择中。')
h('二 玩法与交互创新')
table(['章节','核心玩法','叙事意义'],[
('洪崖洞 迷途','躲避阿姨追逐，穿过流动人群与障碍，触碰石扣开启传送门。','表现主角急躁、只想赶路的初始心态。'),
('磁器口 窑场','在限时压力下，通过移动判定条完成操作流程，逐步打开通路。','体会劳动的难度与生存的不易。'),
('中山古镇 送货','携带货物穿行古镇，应对途中事件，将货物状态带入交付判定。','将诚信与责任落实为需要承担的后果。'),
('防空洞 护送','带领群众前行，利用挡板躲避逐渐加密的轰炸。','从个人逃生转向对他人安全负责。'),
('洪崖洞 归来','自由探索、阅读与交谈，选择拍照或不拍照；照片记录拍摄当天日期。','用相同地点中的不同选择体现人物成长。')],[1.18,3.5,2.22])
p('互动精灵“渝灯”的形象结合灯笼、吊脚楼屋檐与江水意象，位于画面右上方并可拖动。进入关卡后，它会主动介绍相关时代背景，并对点击、拖动做出问候反馈，将操作引导与文化介绍结合起来。')
h('三 AI 工具及具体分工')
table(['工具','实际职责','代表性成果'],[
('DeepSeek Harness','早期开发代理，参与项目框架、场景与基础玩法搭建。','Godot 原型结构、初始关卡与脚本。'),
('OpenAI Codex','参与玩法设计、剧情调整、GDScript 编写、问题排查、资源接入与自动化测试。','五关机制迭代、对话系统、群众跟随、互动精灵、音频管理、转场与发布优化。'),
('Codex 内置 image_gen','OpenAI 图像生成工具，根据角色参考、场景需求和风格约束生成美术，再经筛选与处理接入引擎。','主角动作素材、人物立绘、渝灯、分层背景、互动道具、剧情插画及部分漫画画面。')],[1.45,2.85,2.6])
p('AI 工具记录按实际使用名称披露，不对未完整记录的底层模型版本作额外声明。Godot 是游戏引擎；拆帧、透明化、转码等处理脚本属于制作工具，与生成模型的职责分别说明。')
p('当前正式配乐及部分环境音由创作者提供，片头视频也来自创作者提供的文件；Codex 参与转码、混音、播放衔接和游戏内接入。部分警报、爆炸等声音使用有来源记录的第三方素材。这些内容不统一标注为 AI 原创生成，具体来源以素材清单与授权凭证为准。')
h('四 AI 参与的关键创作过程')
doc.add_heading('玩法设计与脚本实现',2)
p('创作者提出体验目标，例如人群要阻挡主角、群众要随主角跨场景进入第四关、轰炸频率要逐步加快。Codex 将这些目标转化为可运行的脚本、事件判定与参数，并根据试玩反馈持续调整。AI 的贡献落实在正式关卡机制与代码中，可通过源码和提交记录核验。')
doc.add_heading('美术生成与风格控制',2)
p('生成过程设置明确约束，包括角色服装与身份一致、像素密度统一、左右朝向正确、背景远中近层分离、道具透明背景与人物脚底贴地。生成图片经过拆帧、透明化、缩放、锚点校正和实机检查后，才成为正式资源。例如，渝灯采用四帧悬浮循环，主角行走采用六帧循环，背景通过分层产生视差。')
doc.add_heading('叙事组织与人机协作',2)
p('AI 参与补充人物动机与事件因果，使为什么送货、为什么留下救人、传送门如何出现等问题，通过对白、插画和逐格漫画交代，并与关卡行为对应。创作者负责主题、需求、审美判断、试玩反馈和最终取舍；AI 负责大量内容生成、实现与迭代工作，形成由人确定方向、AI 执行并接受验证的协作流程。')
h('五 技术实现与完成度')
p('游戏使用 Godot 场景与 GDScript 组织五关流程，包含玩家移动、NPC 交互、追逐与跟随、剧情状态、关卡重试、音频切换和结局判定。视听表现包括分层视差背景、角色动作、剧情漫画、传送门转场和爆炸临时照明。音乐、环境声与音效分别管理，使不同关卡具有相应的节奏和氛围。')
p('项目已形成包含片头、五关主线和两种结局的 Windows 可运行 Demo。发布阶段通过审计资源、排除未使用的旧素材，将 PCK 从约 140.6 MB 缩减至 79.2 MB，并保留正式关卡与音视频内容。已有回归检查覆盖关卡切换、第四关群众跟随、片头播放、音频路由及剧情资源等环节。')
p('项目保留美术资产清单、生成提示词、部分原始图与处理参数，另有玩法源码、开发记录、Git 提交历史、自动化测试脚本、实机截图与可下载版本，可用于展示 AI 从生成到实际接入的贡献。')
h('六 文化应用价值与使用规范')
p('《洞见》尝试将重庆文化转化为可参与的体验：通过窑场操作感受劳动，通过古镇送货理解承诺，通过防空洞护送体会守望互助。作品可作为城市文化互动展示、研学体验和文旅游戏原型进一步发展。')
p('当前版本的 AI 应用主要发生在开发与内容生产阶段。游戏运行时的追逐、跟随、轰炸与精灵讲解由脚本、状态和预设内容驱动，尚未接入大模型实时推理、自由对话或在线生成关卡。渝灯的现有讲解是按关卡触发的预设内容，不作为实时生成回答的 AI NPC 进行申报。')
p('第三方素材按项目中的来源与许可记录说明；创作者提供的音乐和视频，其来源、生成平台及授权范围应以对应凭证为准。参赛材料应区分生成内容、第三方素材和人工提供内容，保持 AI 使用披露与实际版本一致。')
h('七 评审维度与材料对应')
table(['评分维度','分值','作品说明与验证材料'],[
('玩法与交互创新','20','五关差异化玩法、叙事与操作对应、可拖动互动精灵；通过实机演示验证。'),
('AI 原生程度','25','披露 AI 在内容生产和机制制作中的贡献，提供提示词、资产清单与代码记录；明确当前无运行时模型推理。'),
('技术实现','20','Godot 场景与脚本系统、跨关跟随修复、音频路由和资源包优化；提供源码与测试记录。'),
('完成度与可玩性','15','Windows 可运行 Demo、五关流程、重试和结局；提供完整试玩版与功能演示。'),
('美术与体验设计','10','角色风格控制、分层背景、剧情漫画、临时照明和分章节音频。'),
('应用价值与规范','10','重庆文化互动表达、研学与文旅应用潜力；提供素材来源、授权与 AI 使用说明。')],[1.6,.48,4.82])
p('项目源码：https://github.com/yuanbenkun-cloud/cs2\n试玩发布页：https://github.com/yuanbenkun-cloud/cs2/releases/tag/v1.0.0-demo')
doc.core_properties.title='洞见游戏说明及 AI 应用说明'
doc.core_properties.subject='AI 游戏电竞赛道参赛作品说明'
doc.core_properties.author=''
out=ROOT/'docs'/'洞见_AI游戏电竞赛道_游戏及AI应用说明.docx'
doc.save(out)
print(out)
